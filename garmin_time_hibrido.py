#!/usr/bin/env python3
"""
Time Híbrido -> Garmin Connect
==============================
Cria e agenda os treinos de CORRIDA das 12 semanas do plano Time Híbrido
(Filipe - 5 km, Iniciante Nível 2) direto no calendário do Garmin Connect.

Como funciona
-------------
1. Sem argumentos  -> PRÉVIA OFFLINE: monta os 35 treinos, calcula as datas,
   imprime a agenda e grava `th_preview.json`. NÃO acessa a internet, NÃO
   precisa de login. Use pra conferir se o plano foi lido corretamente.
2. Com --apply     -> Loga no Garmin, ignora treinos já existentes (mesmo nome)
   e cria + agenda cada um.

Credenciais (você define, eu nunca as vejo)
-------------------------------------------
    export GARMIN_EMAIL="seu@email.com"
    export GARMIN_PASSWORD="sua_senha"
    # opcional: onde guardar os tokens (default ~/.garminconnect)
    export GARMINTOKENS="$HOME/.garminconnect"

Uso
---
    python garmin_time_hibrido.py --start 2026-07-06                 # prévia
    python garmin_time_hibrido.py --start 2026-07-06 --apply         # aplica
    python garmin_time_hibrido.py --delete-th --apply                # remove os "S01 · ..."

--start é a SEGUNDA-FEIRA em que a semana 1 começa.
"""
from __future__ import annotations

import argparse
import base64
import contextlib
import json
import os
import re
import shlex
import sys
import time
from datetime import date, datetime, timedelta
from urllib.parse import urlparse

try:
    import montar_musculacao as mm
except Exception:
    mm = None

# IDs oficiais do Garmin workout-service. Valores conferidos contra as
# constantes da biblioteca garminconnect (módulo garminconnect.workout);
# fixados aqui pra o script depender só do garminconnect base.
class SportType:
    RUNNING = 1


class StepType:
    WARMUP, COOLDOWN, INTERVAL, RECOVERY, REST, REPEAT = 1, 2, 3, 4, 5, 6


class ConditionType:
    LAP_BUTTON, TIME, DISTANCE, ITERATIONS = 1, 2, 3, 7


class TargetType:
    NO_TARGET, HEART_RATE_ZONE = 1, 4

REQUEST_DELAY = 1.0  # segundos entre chamadas, pra não estressar a API do Garmin

# Nomenclatura dos treinos (decidida em 12/08/2026): "S01 · Intervalado 1:1".
# O label vem da planilha; a semana entra no nome porque sem ela os nomes
# colidem (ex.: "Intervalado 1:1" é 14, 15 e 16 tiros nas semanas 1, 2 e 3).
# Sem prefixo e sem o dia da semana - o contexto vai na descrição do treino.
NAME_FMT = "S{week:02d} · {label}"            # corrida
NAME_FMT_MUSC = "S{week:02d} {treino} · {grupo}"  # musculação (montar_musculacao)
# Regex que reconhece os treinos criados por este script (usado no --delete-th).
NAME_RE = re.compile(r"^S\d{2} (?:[A-Z] )?· ")

# ----------------------------------------------------------------------------
# Zonas de FC do plano (% da FCmáx) - usadas nas descrições
# ----------------------------------------------------------------------------
ZONE_PCT = {1: "50-59%", 2: "60-69%", 3: "70-79%", 4: "80-89%", 5: "90-95%"}

# Dia da semana -> deslocamento a partir da segunda-feira da semana
DAY_OFFSET = {"2ª": 0, "3ª": 1, "4ª": 2, "5ª": 3, "6ª": 4, "Sáb": 5, "Dom": 6}

# ----------------------------------------------------------------------------
# DSL enxuta pra descrever treinos
# ----------------------------------------------------------------------------
def _min(m: float) -> float:
    return float(m) * 60.0


def wu(minutes: float = 5, zone: int = 1) -> dict:
    return {"kind": "warmup", "sec": _min(minutes), "zone": zone}


def cd(minutes: float = 5, zone: int = 1) -> dict:
    return {"kind": "cooldown", "sec": _min(minutes), "zone": zone}


def work(minutes: float, zone: int = 3) -> dict:
    return {"kind": "interval", "sec": _min(minutes), "zone": zone}


def rec(minutes: float, zone: int = 1) -> dict:
    return {"kind": "recovery", "sec": _min(minutes), "zone": zone}


def work_km(km: float, zone: int = 4) -> dict:
    return {"kind": "interval", "meters": float(km) * 1000.0, "zone": zone}


def rep(reps: int, *steps: dict) -> dict:
    return {"kind": "repeat", "reps": reps, "steps": list(steps)}


# Modelos de treino ----------------------------------------------------------
def intervalado(reps, work_min, rec_min, work_zone=3, rec_zone=1, wu_min=5, cd_min=5):
    return [wu(wu_min), rep(reps, work(work_min, work_zone), rec(rec_min, rec_zone)), cd(cd_min)]


def continuo(run_min, run_zone=3, wu_min=5, cd_min=5):
    return [wu(wu_min), work(run_min, run_zone), cd(cd_min)]


def prova_5km(zone=4, wu_min=5, cd_min=5):
    return [wu(wu_min), work_km(5, zone), cd(cd_min)]


# ----------------------------------------------------------------------------
# O PLANO - 12 semanas (extraído do PDF Time Híbrido)
# (semana, etapa, {dia: (rótulo, passos)})
# Semana 12: 5ª feira é DESCANSO; a prova é no DOMINGO.
# ----------------------------------------------------------------------------
PROGRAMA = [
    (1, "Base", {
        "3ª": ("Intervalado 1:1", intervalado(14, 1, 1)),
        "5ª": ("Intervalado 2:1", intervalado(10, 2, 1)),
        "Sáb": ("Intervalado 3:1", intervalado(7, 3, 1)),
    }),
    (2, "Base", {
        "3ª": ("Intervalado 1:1", intervalado(15, 1, 1)),
        "5ª": ("Intervalado 2:1", intervalado(11, 2, 1)),
        "Sáb": ("Intervalado 3:1", intervalado(8, 3, 1)),
    }),
    (3, "Base", {
        "3ª": ("Intervalado 1:1", intervalado(16, 1, 1)),
        "5ª": ("Intervalado 2:1", intervalado(12, 2, 1)),
        "Sáb": ("Intervalado 3:1", intervalado(9, 3, 1)),
    }),
    (4, "Base", {
        "3ª": ("Intervalado 1:1", intervalado(14, 1, 1)),
        "5ª": ("Intervalado 2:1", intervalado(10, 2, 1)),
        "Sáb": ("Intervalado 3:1", intervalado(7, 3, 1)),
    }),
    (5, "Consolidação", {
        "3ª": ("Intervalado 4:2", intervalado(5, 4, 2)),
        "5ª": ("Intervalado 4:1", intervalado(6, 4, 1)),
        "Sáb": ("Contínuo 20 min", continuo(20)),
    }),
    (6, "Consolidação", {
        "3ª": ("Intervalado 4:2", intervalado(6, 4, 2)),
        "5ª": ("Intervalado 4:1", intervalado(7, 4, 1)),
        "Sáb": ("Contínuo 22 min", continuo(22)),
    }),
    (7, "Consolidação", {
        "3ª": ("Intervalado 4:2", intervalado(7, 4, 2)),
        "5ª": ("Intervalado 4:1", intervalado(8, 4, 1)),
        "Sáb": ("Contínuo 24 min", continuo(24)),
    }),
    (8, "Consolidação", {
        "3ª": ("Intervalado 4:2", intervalado(5, 4, 2)),
        "5ª": ("Intervalado 4:1", intervalado(6, 4, 1)),
        "Sáb": ("Contínuo 20 min", continuo(20)),
    }),
    (9, "Alvo", {
        "3ª": ("Intervalado 5:1", intervalado(6, 5, 1)),
        "5ª": ("Intervalado 6:1", intervalado(6, 6, 1)),
        "Sáb": ("Contínuo 22 min", continuo(22)),
    }),
    (10, "Alvo", {
        "3ª": ("Intervalado 5:1", intervalado(7, 5, 1)),
        "5ª": ("Intervalado 6:1", intervalado(7, 6, 1)),
        "Sáb": ("Contínuo 24 min", continuo(24)),
    }),
    (11, "Alvo", {
        "3ª": ("Intervalado 5:1", intervalado(8, 5, 1)),
        "5ª": ("Intervalado 6:1", intervalado(8, 6, 1)),
        "Sáb": ("Contínuo 26 min", continuo(26)),
    }),
    (12, "Alvo", {
        "3ª": ("Intervalado 5:1", intervalado(6, 5, 1)),
        "Dom": ("PROVA ALVO 5 KM", prova_5km()),
    }),
]

# ----------------------------------------------------------------------------
# Descrição e duração estimada
# ----------------------------------------------------------------------------
_PACE_SECS_PER_KM = 360  # ~6:00/km só pra estimar a duração de trechos por distância


def _fmt_dur(step: dict) -> str:
    if "meters" in step:
        return f"{step['meters'] / 1000:.0f} km"
    return f"{int(step['sec'] // 60)} min"


def _leg_dur_secs(step: dict) -> float:
    if "meters" in step:
        return step["meters"] / 1000 * _PACE_SECS_PER_KM
    return step["sec"]


def describe(steps: list[dict]) -> str:
    labels = {"warmup": "Aquec.", "cooldown": "Desaq.",
              "interval": "", "recovery": "Rec.", "rest": "Desc."}
    parts = []
    zones_used = set()
    for s in steps:
        if s["kind"] == "repeat":
            inner = " / ".join(f"{_fmt_dur(x)} Z{x['zone']}" for x in s["steps"])
            for x in s["steps"]:
                zones_used.add(x["zone"])
            parts.append(f"{s['reps']}x [{inner}]")
        else:
            zones_used.add(s["zone"])
            lbl = labels[s["kind"]]
            seg = f"{_fmt_dur(s)} Z{s['zone']}"
            parts.append(f"{lbl} {seg}".strip())
    legenda = "  |  ".join(f"Z{z} {ZONE_PCT[z]} FCmáx" for z in sorted(zones_used))
    return " · ".join(parts) + "\n" + legenda


def est_secs(steps: list[dict]) -> int:
    total = 0.0
    for s in steps:
        if s["kind"] == "repeat":
            total += sum(_leg_dur_secs(x) for x in s["steps"]) * s["reps"]
        else:
            total += _leg_dur_secs(s)
    return int(total)


# ----------------------------------------------------------------------------
# DSL -> JSON do Garmin (schema real do workout-service)
# ----------------------------------------------------------------------------
_STEP_META = {
    "warmup":   (StepType.WARMUP,   "warmup",   1),
    "cooldown": (StepType.COOLDOWN, "cooldown", 2),
    "interval": (StepType.INTERVAL, "interval", 3),
    "recovery": (StepType.RECOVERY, "recovery", 4),
    "rest":     (StepType.REST,     "rest",     5),
}


def _hr_target(zone: int) -> dict:
    return {
        "workoutTargetTypeId": TargetType.HEART_RATE_ZONE,
        "workoutTargetTypeKey": "heart.rate.zone",
        "displayOrder": 4,
    }


def _end_condition(step: dict):
    if "meters" in step:
        return ({"conditionTypeId": ConditionType.DISTANCE, "conditionTypeKey": "distance",
                 "displayOrder": 3, "displayable": True}, float(step["meters"]))
    return ({"conditionTypeId": ConditionType.TIME, "conditionTypeKey": "time",
             "displayOrder": 2, "displayable": True}, float(step["sec"]))


def _exec_step(step: dict, order: int) -> dict:
    sid, skey, disp = _STEP_META[step["kind"]]
    cond, value = _end_condition(step)
    return {
        "type": "ExecutableStepDTO",
        "stepOrder": order,
        "stepType": {"stepTypeId": sid, "stepTypeKey": skey, "displayOrder": disp},
        "endCondition": cond,
        "endConditionValue": value,
        "targetType": _hr_target(step["zone"]),
        "zoneNumber": step["zone"],
    }


def _build(steps: list[dict], order: int):
    out = []
    for s in steps:
        if s["kind"] == "repeat":
            group_order = order
            order += 1
            child, order = _build(s["steps"], order)
            out.append({
                "type": "RepeatGroupDTO",
                "stepOrder": group_order,
                "stepType": {"stepTypeId": StepType.REPEAT, "stepTypeKey": "repeat", "displayOrder": 6},
                "numberOfIterations": s["reps"],
                "smartRepeat": False,
                "endCondition": {"conditionTypeId": ConditionType.ITERATIONS, "conditionTypeKey": "iterations",
                                 "displayOrder": 7, "displayable": False},
                "endConditionValue": float(s["reps"]),
                "workoutSteps": child,
            })
        else:
            out.append(_exec_step(s, order))
            order += 1
    return out, order


def build_workout_json(name: str, description: str, steps: list[dict]) -> dict:
    workout_steps, _ = _build(steps, 1)
    running = {"sportTypeId": SportType.RUNNING, "sportTypeKey": "running", "displayOrder": 1}
    return {
        "workoutName": name,
        "description": description,
        "sportType": running,
        "estimatedDurationInSecs": est_secs(steps),
        "workoutSegments": [{
            "segmentOrder": 1,
            "sportType": running,
            "workoutSteps": workout_steps,
        }],
    }


# ----------------------------------------------------------------------------
# Montagem da lista completa (treino + data)
# ----------------------------------------------------------------------------
def build_all(start_monday: date, from_week: int = 1, program=None) -> list[dict]:
    prog = program if program is not None else PROGRAMA
    items = []
    for week, stage, days in prog:
        if week < from_week:
            continue
        # --start é a segunda-feira da PRIMEIRA semana que será enviada.
        monday = start_monday + timedelta(days=7 * (week - from_week))
        for day, entry in days.items():
            label, steps = entry[0], entry[1]
            notas = entry[2] if len(entry) > 2 else ""
            when = monday + timedelta(days=DAY_OFFSET[day])
            name = NAME_FMT.format(week=week, label=label)
            desc = f"Time Híbrido · Semana {week:02d} · Etapa {stage}\n{describe(steps)}"
            if notas:
                desc += f"\nNota: {notas}"
            items.append({
                "week": week, "stage": stage, "day": day,
                "date": when.isoformat(), "name": name,
                "payload": build_workout_json(name, desc, steps),
            })
    items.sort(key=lambda x: x["date"])
    return items


def load_program_from_sheet(path: str):
    """Lê o plano de uma planilha .xlsx (aba 'Corrida') e devolve a mesma
    estrutura do PROGRAMA embutido. Cada linha vira um treino."""
    from openpyxl import load_workbook
    if not os.path.exists(path):
        sys.exit(f"ERRO: planilha não encontrada: {path}")
    wb = load_workbook(path, data_only=True)
    ws = wb["Corrida"] if "Corrida" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        sys.exit("ERRO: planilha vazia.")
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    for req in ("semana", "dia", "tipo"):
        if req not in idx:
            sys.exit(f"ERRO: planilha sem a coluna '{req}'. Colunas lidas: {header}")

    def val(row, name, default=None):
        i = idx.get(name)
        if i is None or i >= len(row):
            return default
        v = row[i]
        return default if (v is None or v == "") else v

    program: dict = {}
    for row in rows[1:]:
        if val(row, "semana") is None:
            continue
        week = int(val(row, "semana"))
        stage = str(val(row, "etapa", ""))
        day = str(val(row, "dia")).strip()
        tipo = str(val(row, "tipo", "")).strip().lower()
        aq = int(val(row, "aquecimento_min", 5))
        des = int(val(row, "desaquecimento_min", 5))
        notas = str(val(row, "notas", "") or "")

        if tipo.startswith("interval"):
            reps = int(val(row, "reps"))
            em = float(val(row, "esforco_min"))
            ez = int(val(row, "esforco_zona", 3))
            rm = float(val(row, "recuperacao_min"))
            rz = int(val(row, "recuperacao_zona", 1))
            steps = intervalado(reps, em, rm, work_zone=ez, rec_zone=rz, wu_min=aq, cd_min=des)
            label = f"Intervalado {em:g}:{rm:g}"
        elif tipo.startswith("cont"):
            em = float(val(row, "esforco_min"))
            ez = int(val(row, "esforco_zona", 3))
            steps = continuo(em, run_zone=ez, wu_min=aq, cd_min=des)
            label = f"Contínuo {em:g} min"
        elif tipo.startswith("prova"):
            dist = float(val(row, "distancia_km", 5))
            ez = int(val(row, "esforco_zona", 4))
            steps = [wu(aq), work_km(dist, ez), cd(des)]
            label = f"Prova {dist:g} km"
        else:
            sys.exit(f"ERRO: semana {week} {day}: tipo '{tipo}' desconhecido "
                     "(use intervalado / contínuo / prova).")

        program.setdefault(week, [stage, {}])
        if stage:
            program[week][0] = stage
        program[week][1][day] = (label, steps, notas)

    return [(w, program[w][0], program[w][1]) for w in sorted(program)]


# ----------------------------------------------------------------------------
# Prévia offline
# ----------------------------------------------------------------------------
def preview(items: list[dict]) -> None:
    print(f"\n{'='*66}\nPRÉVIA · {len(items)} treinos de corrida\n{'='*66}")
    cur = None
    for it in items:
        if it["week"] != cur:
            cur = it["week"]
            print(f"\n— Semana {it['week']:02d} ({it['stage']}) —")
        dur = it["payload"]["estimatedDurationInSecs"] // 60
        dt = datetime.fromisoformat(it["date"]).strftime("%a %d/%m")
        print(f"  {dt}  {it['name']:<34} ~{dur:>2} min")
    out = os.path.join(os.getcwd(), "th_preview.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    print(f"\nDetalhe completo (JSON) salvo em: {out}")
    print("Confira e, se estiver certo, rode de novo com --apply.\n")


# ----------------------------------------------------------------------------
# Aplicar no Garmin
# ----------------------------------------------------------------------------
COOLDOWN_FILE = ".login_cooldown"   # marca "não tente logar antes de X" (pós-429)


def _tokenstore() -> str:
    return os.getenv("GARMINTOKENS", os.path.expanduser("~/.garminconnect"))


def _cooldown_file(store: str) -> str:
    return os.path.join(store, COOLDOWN_FILE)


def _check_cooldown(store: str, force: bool) -> None:
    """Impede uma nova tentativa de login logo após um 429 (cada tentativa
    renova o bloqueio da conta). Ignora se --force."""
    if force:
        return
    path = _cooldown_file(store)
    if not os.path.exists(path):
        return
    try:
        until = float(open(path).read().strip())
    except Exception:
        return
    if time.time() < until:
        quando = datetime.fromtimestamp(until).strftime("%H:%M")
        restante = int((until - time.time()) // 60) + 1
        sys.exit(
            f"⏳ Login em espera pós-429 até ~{quando} (faltam ~{restante} min).\n"
            "   Cada tentativa renova o bloqueio da CONTA, então segure a mão.\n"
            "   Se você tem certeza de que já liberou, rode com --force."
        )


def _set_cooldown(store: str, minutes: int) -> None:
    os.makedirs(store, exist_ok=True)
    with contextlib.suppress(Exception):
        with open(_cooldown_file(store), "w") as f:
            f.write(str(time.time() + minutes * 60))


def _clear_cooldown(store: str) -> None:
    with contextlib.suppress(Exception):
        os.remove(_cooldown_file(_tokenstore() if store is None else store))


def _default_browser_headers(cookie: str, csrf: str | None = None) -> dict:
    headers = {
        "Accept": "application/json",
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/126.0.0.0 Safari/537.36"
        ),
        "NK": "NT",
        "Origin": "https://connect.garmin.com",
        "Referer": "https://connect.garmin.com/modern/",
        "DI-Backend": "connectapi.garmin.com",
        "Cookie": cookie,
    }
    if csrf:
        headers["connect-csrf-token"] = csrf.strip()
    return headers


def _headers_from_curl_file(path: str) -> dict:
    text = open(path, encoding="utf-8").read()
    # Chrome/Opera "Copy as cURL (bash)" uses backslash line continuations.
    text = text.replace("\\\r\n", " ").replace("\\\n", " ")
    parts = shlex.split(text, posix=True)
    headers = {}
    curl_url = None
    i = 0
    while i < len(parts):
        part = parts[i]
        value = None
        if part == "curl" and i + 1 < len(parts):
            curl_url = parts[i + 1]
            i += 1
        elif part in ("-H", "--header") and i + 1 < len(parts):
            value = parts[i + 1]
            i += 1
        elif part in ("-b", "--cookie", "--cookie-jar") and i + 1 < len(parts):
            headers["Cookie"] = parts[i + 1]
            i += 1
        elif part.startswith("-H") and len(part) > 2:
            value = part[2:].strip()
        elif part.startswith("--header="):
            value = part.split("=", 1)[1]
        elif part.startswith("--cookie="):
            headers["Cookie"] = part.split("=", 1)[1]
        if value and ":" in value:
            key, val = value.split(":", 1)
            clean_key = key.strip()
            if clean_key.lower() == "cookie":
                clean_key = "Cookie"
            headers[clean_key] = val.strip()
        i += 1

    if "Cookie" not in headers:
        raise RuntimeError("não encontrei header Cookie no arquivo cURL")

    jwt_match = None
    for part in headers["Cookie"].split(";"):
        name, _, value = part.strip().partition("=")
        if name == "JWT_WEB":
            jwt_match = value
            break
    if jwt_match:
        with contextlib.suppress(Exception):
            payload_part = jwt_match.split(".")[1]
            payload = json.loads(
                base64.urlsafe_b64decode(
                    payload_part + "=" * (-len(payload_part) % 4)
                ).decode()
            )
            exp = payload.get("exp")
            if exp and exp < datetime.now().timestamp():
                when = datetime.fromtimestamp(exp).strftime("%Y-%m-%d %H:%M:%S")
                raise RuntimeError(
                    f"a sessão JWT_WEB do garmin-curl.txt expirou em {when}. "
                    "Copie um novo cURL do Opera e substitua o arquivo."
                )

    if curl_url:
        headers["_curl_url"] = curl_url
    headers.setdefault("Accept", "application/json")
    headers.setdefault("Origin", "https://connect.garmin.com")
    headers.setdefault("Referer", "https://connect.garmin.com/modern/")
    headers.setdefault("NK", "NT")
    return headers


def _client_from_web_cookie():
    """Use an already-authenticated Garmin web session as a login fallback.

    Set GARMIN_JWT_WEB from the browser cookie on connect.garmin.com. This avoids
    the fresh SSO login path that can get stuck on 429 or email-code screens.
    Alternatively set GARMIN_COOKIE with the full Cookie request header copied
    from a working browser request to connectapi.garmin.com, or GARMIN_CURL_FILE
    with a DevTools "Copy as cURL (bash)" export for that request.
    """
    curl_file = os.getenv("GARMIN_CURL_FILE")
    jwt_web = os.getenv("GARMIN_JWT_WEB")
    raw_cookie = os.getenv("GARMIN_COOKIE")
    if not curl_file and not jwt_web and not raw_cookie:
        return None

    # JWT_WEB sozinho vira um header Cookie completo. (O caminho antigo chamava
    # garmin._load_profile_and_settings(), método interno que deixou de existir
    # no garminconnect 0.3.x — unificar aqui evita depender dele.)
    if jwt_web and not curl_file and not raw_cookie:
        raw_cookie = f"JWT_WEB={jwt_web.strip()}"

    import garminconnect
    import requests

    garmin = garminconnect.Garmin()
    if jwt_web:
        garmin.client.jwt_web = jwt_web.strip()
    csrf = os.getenv("GARMIN_CSRF")
    if csrf:
        garmin.client.csrf_token = csrf.strip()

    try:
        if curl_file or raw_cookie:
            if curl_file:
                headers = _headers_from_curl_file(curl_file)
                curl_url = headers.pop("_curl_url", None)
                if csrf:
                    headers["connect-csrf-token"] = csrf.strip()
                if curl_url:
                    parsed = urlparse(curl_url)
                    if parsed.netloc == "connect.garmin.com" and parsed.path.startswith("/gc-api/"):
                        garmin.client._connectapi = "https://connect.garmin.com/gc-api"
            else:
                headers = _default_browser_headers(raw_cookie.strip(), csrf)
                # Espelha o navegador: o app web chama connect.garmin.com/gc-api
                # (mesma origem dos cookies). O host connectapi.garmin.com passou
                # a devolver 403 para sessão por cookie (observado em ago/2026).
                curl_url = "https://connect.garmin.com/gc-api/userprofile-service/socialProfile"
                garmin.client._connectapi = "https://connect.garmin.com/gc-api"
            resp = requests.get(
                curl_url or "https://connectapi.garmin.com/userprofile-service/socialProfile",
                headers=headers,
                timeout=20,
            )
            resp.raise_for_status()
            prof = resp.json() if resp.text.strip() else {}
            if not isinstance(prof, dict):   # ex.: /workouts devolve uma lista
                prof = {}
            garmin.client.jwt_web = ""
            garmin.client.get_api_headers = lambda: headers.copy()
            garmin.display_name = prof.get("displayName") or prof.get("userName") or "conta conectada"
            garmin.full_name = prof.get("fullName", "")
            return garmin
    except Exception as e:
        detail = ""
        with contextlib.suppress(Exception):
            if curl_file:
                headers = _headers_from_curl_file(curl_file)
                curl_url = headers.pop("_curl_url", None)
                if csrf:
                    headers["connect-csrf-token"] = csrf.strip()
            else:
                test_cookie = raw_cookie.strip() if raw_cookie else f"JWT_WEB={jwt_web.strip()}"
                headers = _default_browser_headers(test_cookie, csrf)
                curl_url = "https://connect.garmin.com/gc-api/userprofile-service/socialProfile"
            resp = requests.get(
                curl_url or "https://connectapi.garmin.com/userprofile-service/socialProfile",
                headers=headers,
                timeout=20,
            )
            detail = (
                f"\nTeste direto da sessão web: HTTP {resp.status_code}; "
                f"resposta: {resp.text[:240]}"
            )
        source = "GARMIN_CURL_FILE" if curl_file else (
            "GARMIN_COOKIE" if os.getenv("GARMIN_COOKIE") else "GARMIN_JWT_WEB")
        msg = (f"{source} foi informado, mas a sessão web não foi aceita pela API "
               f"(expirada ou incompleta). Erro: {e}{detail}")
        # NÃO cair automaticamente no login por credencial: com o rate limit
        # atual da Garmin (mar/2026), cada tentativa de login renova o bloqueio
        # 429 do IP/conta. Falhou a sessão web -> parar aqui, sem gastar nada.
        sys.exit(
            msg + "\nNADA foi tentado no login por credencial (para não renovar o 429).\n"
            "→ Capture uma sessão nova do navegador: cookie COMPLETO (header\n"
            "  'cookie' inteiro da requisição, em GARMIN_COOKIE) ou Copy as cURL\n"
            "  (bash) em GARMIN_CURL_FILE. Só o JWT_WEB pode não bastar.\n"
            "→ Se você QUER tentar o login por e-mail/senha, rode numa janela\n"
            "  SEM GARMIN_CURL_FILE / GARMIN_COOKIE / GARMIN_JWT_WEB definidos."
        )
    return garmin


def _exigir_stack_de_login() -> None:
    """Aborta ANTES de tocar no SSO se o ambiente não tem o stack anti-429.

    Desde mar/2026 a Garmin bloqueia via Cloudflare (fingerprint TLS + 429) o
    login das libs antigas. O caminho que funciona é o garminconnect >= 0.3.2
    (estratégia "widget+cffi") com o pacote curl_cffi instalado — sem ele a
    estratégia nem é ativada e a tentativa só renova o bloqueio da conta.
    """
    problemas = []
    try:
        import curl_cffi  # noqa: F401
    except ImportError:
        problemas.append("o pacote curl_cffi não está instalado")
    with contextlib.suppress(Exception):
        from importlib.metadata import version
        ver = version("garminconnect")
        partes = tuple(int(p) for p in ver.split(".")[:3])
        if partes < (0, 3, 2):
            problemas.append(f"garminconnect {ver} é antigo (precisa >= 0.3.2)")
    if problemas:
        sys.exit(
            "🛑 Login abortado ANTES de acessar a Garmin: "
            + "; ".join(problemas) + ".\n"
            "   Sem esse stack, a tentativa cai nas estratégias bloqueadas pela\n"
            "   Cloudflare (mudança da Garmin em mar/2026) e só renova o 429.\n"
            "   Corrija com:\n"
            "     py -m pip install -U garminconnect curl_cffi\n"
            "   e rode de novo. (Tokens já salvos continuam funcionando.)"
        )


def get_client(force: bool = False):
    """Autentica priorizando os tokens salvos.

    1) Tenta retomar a sessão pelos tokens salvos — isso NÃO acessa o endpoint
       de login (SSO) e, portanto, é imune ao 429. Não precisa de e-mail/senha.
    2) Se houver GARMIN_JWT_WEB, usa a sessão web já autenticada do navegador.
    3) Só se não houver tokens válidos é que faz o login com credenciais (aí
       sim acessa o SSO e fica sujeito ao 429). Nesse caso exige as variáveis
       GARMIN_EMAIL / GARMIN_PASSWORD.
    """
    import garminconnect
    from garminconnect import GarminConnectTooManyRequestsError

    store = _tokenstore()

    # 1) Retomar via tokens (sem senha, sem SSO). Só tenta se o diretório de
    #    tokens existe — assim não dispara um login vazio (nem SSO) na 1ª vez.
    if os.path.isdir(store) and os.listdir(store):
        try:
            garmin = garminconnect.Garmin()
            garmin.login(store)
            _clear_cooldown(store)
            return garmin
        except Exception:
            pass  # tokens inválidos/expirados -> login com credencial abaixo

    # 2) Reaproveitar sessão web já logada (sem senha, sem SSO)
    garmin = _client_from_web_cookie()
    if garmin is not None:
        _clear_cooldown(store)
        return garmin

    # 3) Login com credencial (primeira vez ou tokens expirados de vez)
    email = os.getenv("GARMIN_EMAIL")
    password = os.getenv("GARMIN_PASSWORD")
    if not email or not password:
        sys.exit(
            "Primeira autenticação (ainda não há tokens salvos): defina\n"
            "  $env:GARMIN_EMAIL e $env:GARMIN_PASSWORD  e rode de novo.\n"
            "Se o login normal está travado em 429/código, você também pode usar\n"
            "  $env:GARMIN_CURL_FILE apontando para um Copy as cURL do DevTools,\n"
            "  ou $env:GARMIN_JWT_WEB com o cookie JWT_WEB do navegador já logado.\n"
            "Depois deste primeiro login, os tokens ficam salvos e as próximas\n"
            "execuções não pedem mais credenciais nem acessam o login da Garmin."
        )
    _exigir_stack_de_login()
    _check_cooldown(store, force)

    print(
        "→ Login na Garmin (1 tentativa). A estratégia anti-429 pode esperar\n"
        "  30–45 s de propósito antes de enviar a credencial — não interrompa."
    )
    garmin = garminconnect.Garmin(
        email=email, password=password,
        prompt_mfa=lambda: input("Código (só se pedir): ").strip(),
    )
    try:
        garmin.login(store)          # em caso de sucesso, salva os tokens sozinho
    except GarminConnectTooManyRequestsError:
        _set_cooldown(store, 60)
        sys.exit(
            "🚫 429 — a Garmin recusou o login com 'Too Many Requests'. É disparado\n"
            "   por tentativas seguidas de login. Pode ser bloqueio por IP OU pela conta:\n"
            "   • De OUTRA rede (ex.: hotspot do celular) às vezes libera na hora;\n"
            "     se persistir mesmo assim, é pela conta e leva de ~1h a alguns dias.\n"
            "   • PARE de tentar — cada tentativa renova o bloqueio.\n"
            "   • Alternativa imediata: sessão do navegador (GARMIN_CURL_FILE), que\n"
            "     não passa pelo login e ignora o 429.\n"
            "   • Quando passar, rode UMA vez; os tokens ficam salvos e não volta.\n"
            "   (Trava local de 1h ativada; use --force para pular.)"
        )
    except Exception as e:
        # NÃO chame tudo de 429: diagnostique a causa real (isso enganava em rede
        # corporativa, onde uma falha de SSL era reportada como 'limite de login').
        msg = str(e)
        low = msg.lower()
        if any(k in low for k in ("certificate", "ssl", "certificate_verify_failed", "self signed")):
            sys.exit(
                "🔒 Falha de SSL/certificado ao acessar o login da Garmin — NÃO é 429.\n"
                "   Típico de rede corporativa/antivírus/VPN que faz inspeção de SSL.\n"
                "   Correção no Windows (usa a loja de certificados do sistema):\n"
                "     py -m pip install pip-system-certs\n"
                "   ou aponte o certificado raiz da empresa:\n"
                "     $env:REQUESTS_CA_BUNDLE = 'C:\\caminho\\empresa-ca.pem'\n"
                "     $env:CURL_CA_BUNDLE     = 'C:\\caminho\\empresa-ca.pem'\n"
                f"   Erro original: {msg}"
            )
        if any(k in low for k in ("429", "too many", "rate limit")):
            _set_cooldown(store, 60)
            sys.exit("🚫 429 — limite de login. Pare ~1h e tente UMA vez "
                     "(ou use hotspot / GARMIN_CURL_FILE).")
        # Qualquer outra causa (senha errada, 401, rede): mostre o erro REAL.
        sys.exit(f"✗ Login falhou (não foi 429): {msg}")
    _clear_cooldown(store)
    return garmin


def check_login(force: bool = False) -> None:
    """Só verifica o login (uma requisição) e sai. Útil pra testar se o 429
    passou, sem disparar o upload inteiro."""
    garmin = get_client(force=force)
    quem = getattr(garmin, "full_name", None) or getattr(garmin, "display_name", None) or "sua conta"
    print(f"✓ Login OK — sessão ativa para: {quem}")
    print("Tokens salvos. As próximas execuções não vão logar de novo nem acessar o SSO.")


def apply(items: list[dict], force: bool = False) -> None:
    garmin = get_client(force=force)

    existing = {}
    try:
        for w in garmin.get_workouts(0, 999):
            existing[w.get("workoutName")] = w.get("workoutId")
    except Exception as e:
        print(f"Aviso: não consegui listar treinos existentes ({e}). Sigo mesmo assim.")

    ja_tem = sum(1 for it in items if it["name"] in existing)
    faltam = len(items) - ja_tem
    print(f"Situação: {ja_tem} de {len(items)} já no Garmin · {faltam} a criar.\n")

    created = skipped = scheduled = failed = 0
    for it in items:
        name, when = it["name"], it["date"]
        if name in existing:
            print(f"• pulei (já existe): {name}")
            skipped += 1
            continue
        try:
            res = garmin.upload_workout(it["payload"])
            wid = res.get("workoutId") if isinstance(res, dict) else None
            if not wid:
                raise RuntimeError(f"resposta sem workoutId: {res!r}")
            created += 1
            time.sleep(REQUEST_DELAY)
            garmin.schedule_workout(wid, when)
            scheduled += 1
            print(f"✓ criado + agendado {when}: {name}  (id {wid})")
            time.sleep(REQUEST_DELAY)
        except Exception as e:
            failed += 1
            print(f"✗ FALHOU: {name} -> {e}")
            if any(k in str(e).lower() for k in ("429", "rate", "too many")):
                print("→ Parando: limite de requisições atingido. Rode de novo mais "
                      "tarde — os que já subiram serão pulados automaticamente.")
                break

    print(f"\nResumo: {created} criados · {scheduled} agendados · "
          f"{skipped} já existiam · {failed} falharam")
    if created:
        print("Abra o Garmin Connect › Calendário e sincronize o relógio. 🏃")


def delete_th(force: bool = False) -> None:
    garmin = get_client(force=force)
    # Reconhece os nomes deste script ("S01 · ..." e "S01 A · ...") e também os
    # do formato antigo ("TH S01 3ª · ...", "TH-M S01 A · ..."), pra limpar
    # sobras de execuções anteriores à mudança de nomenclatura.
    def meu(nome: str) -> bool:
        return bool(NAME_RE.match(nome)) or nome.startswith(("TH S", "TH-M S"))

    alvo = [w for w in garmin.get_workouts(0, 999)
            if meu(str(w.get("workoutName", "")))]
    if not alvo:
        print("Nenhum treino do Time Híbrido encontrado.")
        return
    print(f"Vou APAGAR {len(alvo)} treinos do Time Híbrido:")
    for w in alvo:
        print("   -", w.get("workoutName"))
    if input("Digite APAGAR para confirmar: ").strip() != "APAGAR":
        print("Cancelado.")
        return
    for w in alvo:
        try:
            garmin.delete_workout(w["workoutId"])
            print("apagado:", w.get("workoutName"))
            time.sleep(REQUEST_DELAY)
        except Exception as e:
            print("erro ao apagar", w.get("workoutName"), "->", e)


# ----------------------------------------------------------------------------
def parse_start(value: str | None) -> date:
    value = value or os.getenv("TH_START_MONDAY")
    if not value:
        sys.exit("ERRO: informe a segunda-feira de início com --start AAAA-MM-DD "
                 "(ou a variável TH_START_MONDAY).")
    try:
        d = datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        sys.exit("ERRO: data inválida. Use o formato AAAA-MM-DD, ex.: 2026-07-06.")
    if d.weekday() != 0:  # 0 = segunda
        nome = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"][d.weekday()]
        print(f"AVISO: {value} é {nome}, não uma segunda-feira. "
              "As datas dos treinos serão contadas a partir dessa data mesmo assim.")
    return d


def _walk(steps):
    for s in steps:
        if s.get("type") == "RepeatGroupDTO":
            yield from _walk(s["workoutSteps"])
        else:
            yield s


def preview_strength(workouts) -> None:
    print(f"\n{'='*62}\nPRÉVIA MUSCULAÇÃO · {len(workouts)} treinos\n{'='*62}")
    cur = None
    for w in workouts:
        if w["week"] != cur:
            cur = w["week"]
            print(f"\n— Semana {w['week']:02d} —")
        n_ex = sum(1 for s in _walk(w["payload"]["workoutSegments"][0]["workoutSteps"])
                   if s.get("exerciseName"))
        print(f"  {w['name']:<40} {n_ex} exercícios")
    print("\nNão publica nada. Rode com --apply para criar no Garmin.")


def apply_strength(workouts, force: bool = False) -> None:
    garmin = get_client(force=force)
    existing = set()
    try:
        for w in garmin.get_workouts(0, 999):
            existing.add(w.get("workoutName"))
    except Exception as e:
        print(f"Aviso: não listei treinos existentes ({e}). Sigo mesmo assim.")

    created = skipped = failed = 0
    for w in workouts:
        if w["name"] in existing:
            print(f"• pulei (já existe): {w['name']}")
            skipped += 1
            continue
        try:
            res = garmin.upload_workout(w["payload"])
            wid = res.get("workoutId") if isinstance(res, dict) else None
            if not wid:
                raise RuntimeError(f"resposta sem workoutId: {res!r}")
            created += 1
            print(f"✓ criado: {w['name']}  (id {wid})")
            time.sleep(REQUEST_DELAY)
        except Exception as e:
            failed += 1
            print(f"✗ FALHOU: {w['name']} -> {e}")
            if any(k in str(e).lower() for k in ("429", "rate", "too many")):
                print("→ Parando: limite de requisições. Rode de novo mais tarde "
                      "(os já criados serão pulados).")
                break

    print(f"\nResumo musculação: {created} criados · {skipped} já existiam · {failed} falharam")
    if created:
        print("Os treinos ficam na sua biblioteca (Garmin Connect › Treinos). Agende-os "
              "ou escolha o Treino A–E no relógio a cada dia de academia. 💪")


def main() -> None:
    p = argparse.ArgumentParser(description="Time Híbrido -> Garmin Connect (corrida).")
    p.add_argument("--start",
                   help="Segunda-feira da 1ª semana que será enviada (AAAA-MM-DD).")
    p.add_argument("--from-week", type=int, default=1, metavar="N",
                   help="Enviar a partir desta semana (1–12). Padrão: 1.")
    p.add_argument("--sheet", metavar="ARQ.xlsx",
                   help="Lê o plano de uma planilha .xlsx (aba 'Corrida') em vez "
                        "do plano embutido.")
    p.add_argument("--apply", action="store_true",
                   help="Realmente cria e agenda no Garmin (sem isso, só prévia).")
    p.add_argument("--check", action="store_true",
                   help="Só testa o login (1 requisição) e sai. Bom p/ ver se o 429 passou.")
    p.add_argument("--force", action="store_true",
                   help="Ignora a trava local de espera pós-429.")
    p.add_argument("--delete-th", action="store_true",
                   help="Apaga os treinos do Time Híbrido - 'S01 · ...' e o "
                        "formato antigo 'TH ...' (requer --apply).")
    p.add_argument("--musculacao", action="store_true",
                   help="Cria os treinos de MUSCULAÇÃO da planilha (requer --sheet).")
    args = p.parse_args()

    if args.check:
        check_login(force=args.force)
        return

    if args.delete_th:
        if not args.apply:
            sys.exit("Por segurança, --delete-th exige também --apply.")
        delete_th(force=args.force)
        return

    if args.musculacao:
        if not args.sheet:
            sys.exit("--musculacao exige --sheet com a planilha (.xlsx).")
        if mm is None:
            sys.exit("Módulo montar_musculacao.py não encontrado (coloque-o na mesma pasta).")
        workouts = mm.build_all_strength(args.sheet)
        if args.apply:
            apply_strength(workouts, force=args.force)
        else:
            preview_strength(workouts)
        return

    if not 1 <= args.from_week <= 12:
        sys.exit("ERRO: --from-week deve estar entre 1 e 12.")
    program = load_program_from_sheet(args.sheet) if args.sheet else None
    start = parse_start(args.start)
    items = build_all(start, args.from_week, program=program)

    if args.apply:
        apply(items, force=args.force)
    else:
        preview(items)


if __name__ == "__main__":
    main()
